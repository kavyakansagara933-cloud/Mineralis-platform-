from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any
import pypdf
import pdfplumber
import openpyxl
from docx import Document as DocxDocument
from PIL import Image
import pytesseract

class DocumentParserService:
    @staticmethod
    def parse_file(file_path: Path, filename: str) -> dict[str, Any]:
        ext = file_path.suffix.lower()
        
        if ext == ".pdf":
            return DocumentParserService._parse_pdf(file_path)
        elif ext in [".xlsx", ".xls"]:
            return DocumentParserService._parse_excel(file_path)
        elif ext == ".csv":
            return DocumentParserService._parse_csv(file_path)
        elif ext == ".docx":
            return DocumentParserService._parse_docx(file_path)
        elif ext in [".png", ".jpg", ".jpeg"]:
            return DocumentParserService._parse_image(file_path)
        else:
            # Fallback text
            try:
                text = file_path.read_text(encoding="utf-8", errors="ignore")
                return {
                    "total_pages": 1,
                    "pages": [{"page_number": 1, "text": text, "tables": [], "ocr_applied": False}],
                    "file_type": "text",
                    "confidence": 85
                }
            except Exception:
                return {
                    "total_pages": 1,
                    "pages": [{"page_number": 1, "text": "Unrecognized format", "tables": [], "ocr_applied": False}],
                    "file_type": "unknown",
                    "confidence": 50
                }

    @staticmethod
    def _parse_pdf(file_path: Path) -> dict[str, Any]:
        pages_data = []
        is_scanned = False
        overall_confidence = 96

        try:
            with pdfplumber.open(file_path) as pdf:
                total_pages = len(pdf.pages)
                for idx, page in enumerate(pdf.pages, start=1):
                    text = page.extract_text() or ""
                    tables = page.extract_tables() or []
                    ocr_applied = False

                    # If page has almost no selectable text, attempt OCR
                    if len(text.strip()) < 40:
                        try:
                            # Attempt OCR on page image
                            page_img = page.to_image(resolution=200).original
                            ocr_text = pytesseract.image_to_string(page_img)
                            if len(ocr_text.strip()) > len(text.strip()):
                                text = ocr_text
                                ocr_applied = True
                                is_scanned = True
                                overall_confidence = 74
                        except Exception:
                            pass

                    pages_data.append({
                        "page_number": idx,
                        "text": text,
                        "tables": tables,
                        "ocr_applied": ocr_applied
                    })
        except Exception:
            # Fallback to pypdf
            reader = pypdf.PdfReader(str(file_path))
            total_pages = len(reader.pages)
            for idx, page in enumerate(reader.pages, start=1):
                text = page.extract_text() or ""
                pages_data.append({
                    "page_number": idx,
                    "text": text,
                    "tables": [],
                    "ocr_applied": False
                })

        return {
            "total_pages": total_pages,
            "pages": pages_data,
            "file_type": "scanned_pdf" if is_scanned else "pdf",
            "confidence": overall_confidence
        }

    @staticmethod
    def _parse_excel(file_path: Path) -> dict[str, Any]:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        pages_data = []
        
        for idx, sheet_name in enumerate(wb.sheetnames, start=1):
            sheet = wb[sheet_name]
            rows = []
            text_lines = [f"Sheet: {sheet_name}"]
            for row in sheet.iter_rows(values_only=True):
                row_vals = [str(v) if v is not None else "" for v in row]
                if any(row_vals):
                    rows.append(row_vals)
                    text_lines.append(" | ".join(row_vals))

            pages_data.append({
                "page_number": idx,
                "text": "\n".join(text_lines),
                "tables": [rows] if rows else [],
                "ocr_applied": False,
                "sheet_name": sheet_name
            })

        return {
            "total_pages": len(pages_data),
            "pages": pages_data,
            "file_type": "xlsx",
            "confidence": 98
        }

    @staticmethod
    def _parse_csv(file_path: Path) -> dict[str, Any]:
        rows = []
        text_lines = []
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            reader = csv.reader(f)
            for row in reader:
                if any(row):
                    rows.append(row)
                    text_lines.append(" | ".join(row))

        return {
            "total_pages": 1,
            "pages": [{
                "page_number": 1,
                "text": "\n".join(text_lines),
                "tables": [rows],
                "ocr_applied": False
            }],
            "file_type": "csv",
            "confidence": 98
        }

    @staticmethod
    def _parse_docx(file_path: Path) -> dict[str, Any]:
        doc = DocxDocument(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        tables_data = []
        for table in doc.tables:
            table_rows = []
            for row in table.rows:
                table_rows.append([cell.text.strip() for cell in row.cells])
            tables_data.append(table_rows)

        full_text = "\n".join(paragraphs)
        return {
            "total_pages": 1,
            "pages": [{
                "page_number": 1,
                "text": full_text,
                "tables": tables_data,
                "ocr_applied": False
            }],
            "file_type": "docx",
            "confidence": 95
        }

    @staticmethod
    def _parse_image(file_path: Path) -> dict[str, Any]:
        ocr_text = ""
        confidence = 68
        try:
            img = Image.open(file_path)
            ocr_text = pytesseract.image_to_string(img)
            confidence = 78 if len(ocr_text.strip()) > 30 else 55
        except Exception:
            ocr_text = "[OCR processing failed or image text unclear]"

        return {
            "total_pages": 1,
            "pages": [{
                "page_number": 1,
                "text": ocr_text,
                "tables": [],
                "ocr_applied": True
            }],
            "file_type": "image",
            "confidence": confidence
        }
