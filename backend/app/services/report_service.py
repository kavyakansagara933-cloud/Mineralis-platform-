from __future__ import annotations

from pathlib import Path
from typing import Any
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from docx import Document as DocxDocument
from ..config import EXPORTS_DIR

class ReportExportService:
    @staticmethod
    def generate_pdf(report_id: str, report_data: dict[str, Any]) -> str:
        pdf_filename = f"{report_id}.pdf"
        pdf_path = EXPORTS_DIR / pdf_filename
        
        doc = SimpleDocTemplate(str(pdf_path), pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
        styles = getSampleStyleSheet()
        story = []

        # Title & Meta
        title_style = ParagraphStyle(
            'ReportTitle',
            parent=styles['Heading1'],
            fontSize=20,
            leading=24,
            textColor=colors.HexColor("#0f172a"),
            spaceAfter=6
        )
        sub_style = ParagraphStyle(
            'ReportSub',
            parent=styles['Normal'],
            fontSize=10,
            textColor=colors.HexColor("#64748b"),
            spaceAfter=14
        )
        h2_style = ParagraphStyle(
            'ReportH2',
            parent=styles['Heading2'],
            fontSize=13,
            leading=16,
            textColor=colors.HexColor("#1e293b"),
            spaceBefore=12,
            spaceAfter=6
        )
        body_style = ParagraphStyle(
            'ReportBody',
            parent=styles['Normal'],
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor("#334155")
        )

        story.append(Paragraph(report_data.get("title", "Mining Intelligence Verified Report"), title_style))
        meta_line = f"Organization: {report_data.get('organization', 'Coal India Limited')} | Period: {report_data.get('period', 'FY 2024-25')} | Status: VERIFIED AUDIT"
        story.append(Paragraph(meta_line, sub_style))
        story.append(Spacer(1, 10))

        # Executive Summary
        story.append(Paragraph("1. Executive Summary", h2_style))
        story.append(Paragraph(report_data.get("executive_summary", "Operational analysis completed with verified mathematical audit trail."), body_style))
        story.append(Spacer(1, 10))

        # Key Findings
        story.append(Paragraph("2. Key Operational Findings", h2_style))
        for finding in report_data.get("key_findings", []):
            story.append(Paragraph(f"• {finding}", body_style))
        story.append(Spacer(1, 10))

        # Metrics Table
        story.append(Paragraph("3. Audited Production & Dispatch Table", h2_style))
        table_data = [["Mine / Asset", "Subsidiary", "Actual (MT)", "Target (MT)", "Variance (%)", "Confidence"]]
        for m in report_data.get("metrics_data", []):
            var_str = f"{m.get('variance_percent', 0.0)}%" if m.get('variance_percent') is not None else "N/A"
            table_data.append([
                str(m.get("mine_name", "Mine")),
                str(m.get("subsidiary", "CIL")),
                f"{m.get('actual_value', 0.0)} {m.get('unit', 'MT')}",
                f"{m.get('target_value', 'N/A')}",
                var_str,
                f"{m.get('confidence', 95)}%"
            ])

        t = Table(table_data, colWidths=[110, 80, 85, 85, 80, 80])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1e293b")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
            ('TOPPADDING', (0, 0), (-1, 0), 6),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ('FONTSIZE', (0, 1), (-1, -1), 8.5),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ]))
        story.append(t)
        story.append(Spacer(1, 14))

        # Evidence Trail
        story.append(Paragraph("4. Traceable Source Provenance & Evidence Audit Trail", h2_style))
        for ev in report_data.get("evidence_data", []):
            ev_line = f"<b>[{ev.get('document_name', 'Doc')}, Page {ev.get('page_number', 1)}]</b> {ev.get('snippet', '')} <i>(Confidence: {ev.get('confidence', 95)}%)</i>"
            story.append(Paragraph(ev_line, body_style))
            story.append(Spacer(1, 4))

        doc.build(story)
        return f"/exports/{pdf_filename}"

    @staticmethod
    def generate_docx(report_id: str, report_data: dict[str, Any]) -> str:
        docx_filename = f"{report_id}.docx"
        docx_path = EXPORTS_DIR / docx_filename

        doc = DocxDocument()
        doc.add_heading(report_data.get("title", "Mining Intelligence Verified Report"), level=0)
        
        meta = f"Organization: {report_data.get('organization', 'Coal India Limited')} | Period: {report_data.get('period', 'FY 2024-25')} | Status: VERIFIED AUDIT"
        doc.add_paragraph(meta)

        doc.add_heading("1. Executive Summary", level=1)
        doc.add_paragraph(report_data.get("executive_summary", ""))

        doc.add_heading("2. Key Operational Findings", level=1)
        for finding in report_data.get("key_findings", []):
            doc.add_paragraph(f"• {finding}")

        doc.add_heading("3. Audited Metrics", level=1)
        table = doc.add_table(rows=1, cols=6)
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Mine / Asset"
        hdr_cells[1].text = "Subsidiary"
        hdr_cells[2].text = "Actual"
        hdr_cells[3].text = "Target"
        hdr_cells[4].text = "Variance"
        hdr_cells[5].text = "Confidence"

        for m in report_data.get("metrics_data", []):
            row_cells = table.add_row().cells
            row_cells[0].text = str(m.get("mine_name", "Mine"))
            row_cells[1].text = str(m.get("subsidiary", "CIL"))
            row_cells[2].text = f"{m.get('actual_value', 0.0)} {m.get('unit', 'MT')}"
            row_cells[3].text = str(m.get('target_value', 'N/A'))
            row_cells[4].text = f"{m.get('variance_percent', 0.0)}%" if m.get('variance_percent') is not None else "N/A"
            row_cells[5].text = f"{m.get('confidence', 95)}%"

        doc.add_heading("4. Evidence Audit Trail", level=1)
        for ev in report_data.get("evidence_data", []):
            p = doc.add_paragraph()
            p.add_run(f"[{ev.get('document_name', 'Doc')}, Page {ev.get('page_number', 1)}] ").bold = True
            p.add_run(ev.get('snippet', ''))

        doc.save(str(docx_path))
        return f"/exports/{docx_filename}"
